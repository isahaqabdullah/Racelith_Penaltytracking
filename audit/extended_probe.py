"""Additional probes; run only against the disposable audit database."""
import json, os, urllib.request, urllib.error, urllib.parse, io, threading, concurrent.futures
from pathlib import Path
from datetime import datetime, timezone, timedelta
import psycopg2
from openpyxl import load_workbook
OUT=Path(__file__).parent
if os.environ.get('RACELITH_AUDIT_DISPOSABLE') != '1':
    raise SystemExit('Set RACELITH_AUDIT_DISPOSABLE=1 only for the disposable audit instance.')
results=[]
def req(method,path,data=None,raw=False):
    try:
        with urllib.request.urlopen(urllib.request.Request('http://127.0.0.1:8000'+path,json.dumps(data).encode() if data is not None else None,{'Content-Type':'application/json'},method=method),timeout=30) as r:
            body=r.read();return r.status,body if raw else json.loads(body)
    except urllib.error.HTTPError as e:
        b=e.read()
        try:b=json.loads(b)
        except:b=b.decode()
        return e.code,b
def record(name,ok,evidence):
    result=dict(name=name,status='PASS' if ok else 'FAIL',evidence=evidence);results.append(result);print(json.dumps(result),flush=True)
def create(kart,**kw):return req('POST','/infringements/',dict(kart_number=kart,description='White Line Infringement',penalty_description='Warning',**kw))
def upload(filename,content):
    boundary='auditMultipart';body=(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n\r\n'.encode()+content+f'\r\n--{boundary}--\r\n'.encode())
    try:
        with urllib.request.urlopen(urllib.request.Request('http://127.0.0.1:8000/session/import',body,{'Content-Type':f'multipart/form-data; boundary={boundary}'},method='POST')) as r:return r.status,json.loads(r.read())
    except urllib.error.HTTPError as e:return e.code,json.loads(e.read())

req('POST','/session/start?name=Audit%20Extended')
expired=datetime.now(timezone.utc)-timedelta(minutes=181)
create(81,timestamp=expired.isoformat()); current=create(81)[1]
record('Expired warnings excluded',current['warning_count']==1,current)
create(82);yellow=req('POST','/infringements/',dict(kart_number=82,description='Yellow Zone Infringement',penalty_description='Warning'))[1]
record('White and yellow warning cycles independent',yellow['warning_count']==1,yellow)
same=datetime.now(timezone.utc).isoformat()
equal=[create(83,timestamp=same)[1] for _ in range(4)]
record('Equal timestamps still reset warning cycle',equal[-1]['warning_count']==1,[(x['warning_count'],x['penalty_due']) for x in equal])
old=datetime.now(timezone.utc)-timedelta(hours=1)
create(84,timestamp=old.isoformat());later=create(84)[1];backdated=create(84,timestamp=(old+timedelta(minutes=1)).isoformat())[1]
record('Backdated warnings exclude later events',backdated['warning_count']==2,{'later':later,'backdated':backdated})
create(85);create(85)
barrier=threading.Barrier(5)
def concurrent_warning(i):barrier.wait();return create(85)[1]
with concurrent.futures.ThreadPoolExecutor(max_workers=5) as pool:
    warnings=list(pool.map(concurrent_warning,range(5)))
record('Concurrent warnings follow one serialized cycle',sorted(x['warning_count'] for x in warnings)==[1,1,2,3,3],[(x.get('id'),x.get('warning_count'),x.get('penalty_due')) for x in warnings])
blank=req('POST','/infringements/',{'kart_number':86})[1]
s,xlsx=req('GET','/session/export?name=Audit%20Extended&format=excel',raw=True)
wb=load_workbook(io.BytesIO(xlsx));wb['Infringements']['B2']='Audit Blank Import';buf=io.BytesIO();wb.save(buf)
s,b=upload('blank.xlsx',buf.getvalue())
record('Blank-description records survive Excel round trip',any(x['kart_number']==86 for x in req('GET','/infringements/')[1]['items']),{'status':s,'originalId':blank['id'],'import':b})
req('POST','/session/load?name=Audit%20Extended')
wb=load_workbook(io.BytesIO(xlsx));ws=wb['Infringements'];ws['B2']='Audit Failed Import'
for row in ws:
    if row[0].value=='ID': header=row[0].row;break
for index in range(header+1,ws.max_row+1):
    if ws.cell(index,4).value:
        ws.cell(index,2,2147483648)
        break
buf=io.BytesIO();wb.save(buf)
s,b=upload('invalid.xlsx',buf.getvalue())
active=[x for x in req('GET','/session/')[1]['sessions'] if x['status']=='active']
record('Failed import preserves prior active session',s>=400 and active[0]['name']=='Audit Extended',{'status':s,'response':b,'active':active})
req('POST','/session/load?name=Audit%20Extended')
bulk=[req('POST','/infringements/',dict(kart_number=87,description='Contact',penalty_description='10 Sec'))[1] for _ in range(2)]
req('POST','/penalties/apply/87',{'performed_by':'Audit'})
history=req('GET','/history/87')[1]
record('Apply-all audits each applied penalty',len([h for h in history if h['action']=='penalty_applied'])==2,history)
# Verify control DB fallback after the deletion of a database by a name alias.
req('POST','/session/start?name=Audit%20Broken')
req('DELETE','/session/delete?name=AUDIT%20BROKEN')
s,b=create(88)
conn=psycopg2.connect('postgresql://audit_user@127.0.0.1:55439/audit_control');cur=conn.cursor();cur.execute('SELECT kart_number FROM infringements WHERE kart_number=88');control_rows=cur.fetchall();cur.close();conn.close()
record('Failed session switch never falls back to control DB writes',s!=200 and not control_rows,{'status':s,'response':b,'controlRows':control_rows})
req('POST','/session/load?name=Audit%20Extended')
(OUT/'extended-results.json').write_text(json.dumps(results,indent=2))
