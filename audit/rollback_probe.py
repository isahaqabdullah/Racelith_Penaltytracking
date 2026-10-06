"""Prove import failure handling against a disposable Racelith instance."""
import os, io, json, urllib.request, urllib.error
from pathlib import Path
from openpyxl import load_workbook
if os.environ.get('RACELITH_AUDIT_DISPOSABLE') != '1':
    raise SystemExit('Set RACELITH_AUDIT_DISPOSABLE=1 only for the disposable audit instance.')
def request(method,path,data=None,raw=False):
    try:
        with urllib.request.urlopen(urllib.request.Request('http://127.0.0.1:8000'+path,json.dumps(data).encode() if data else None,{'Content-Type':'application/json'},method=method)) as r:
            b=r.read();return r.status,b if raw else json.loads(b)
    except urllib.error.HTTPError as e:return e.code,json.loads(e.read())
request('POST','/session/start?name=Audit%20Rollback')
request('POST','/infringements/',{'kart_number':91,'description':'Contact','penalty_description':'10 Sec'})
_,data=request('GET','/session/export?name=Audit%20Rollback&format=excel',raw=True)
wb=load_workbook(io.BytesIO(data));ws=wb['Infringements'];ws['B2']='Audit Bad Rollback'
for row in ws:
    if isinstance(row[0].value,int): ws.cell(row[0].row,2,2147483648);break
buf=io.BytesIO();wb.save(buf)
boundary='auditFailureBoundary';body=(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="bad.xlsx"\r\n\r\n'.encode()+buf.getvalue()+f'\r\n--{boundary}--\r\n'.encode())
try:
    with urllib.request.urlopen(urllib.request.Request('http://127.0.0.1:8000/session/import',body,{'Content-Type':f'multipart/form-data; boundary={boundary}'},method='POST')) as r:status=r.status;response=json.loads(r.read())
except urllib.error.HTTPError as e:status=e.code;response=json.loads(e.read())
active=[s for s in request('GET','/session/')[1]['sessions'] if s['status']=='active']
results=[{'name':'Failed import preserves prior active session','status':'PASS' if status>=400 and active[0]['name']=='Audit Rollback' else 'FAIL','evidence':{'status':status,'response':response,'active':active,'remainingRecords':request('GET','/infringements/')[1]}}]
Path(__file__).with_name('rollback-results.json').write_text(json.dumps(results,indent=2))
print(json.dumps(results,indent=2))
